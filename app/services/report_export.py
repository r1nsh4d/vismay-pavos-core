import io
from datetime import datetime, timezone

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.units import cm, mm
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_RIGHT, TA_LEFT, TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable,
)


# ── Excel ──────────────────────────────────────────────────────────────────────
'''
def generate_excel(rows: list[dict], sheet_name: str = "Report") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # Excel sheet name max 31 chars

    if not rows:
        ws.append(["No data available"])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    headers = list(rows[0].keys())

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Write headers
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Write data rows
    for row_idx, row in enumerate(rows, 2):
        fill = PatternFill(
            start_color="EBF3FB" if row_idx % 2 == 0 else "FFFFFF",
            end_color="EBF3FB" if row_idx % 2 == 0 else "FFFFFF",
            fill_type="solid",
        )
        for col_idx, header in enumerate(headers, 1):
            value = row.get(header, "")
            # Convert None to empty string
            if value is None:
                value = ""
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = thin_border
            cell.fill = fill

    # Auto column width
    for col in ws.columns:
        max_length = max(
            (
                max((len(line) for line in str(cell.value or "").split("\n")), default=0)
                for cell in col
            ),
            default=10,
        )
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 45)

    # Row height for header
    ws.row_dimensions[1].height = 25

    # Freeze header row
    ws.freeze_panes = "A2"

    # Add generated timestamp in a metadata sheet
    meta_ws = wb.create_sheet("Info")
    meta_ws.append(["Generated At", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    meta_ws.append(["Total Rows", len(rows)])
    meta_ws.append(["Sheet", sheet_name])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
'''

def _merge_runs(ws, merge_key: str, merge_cols, header_row: int = 1) -> None:
    """Vertically merge `merge_cols` across contiguous rows sharing the same
    value in `merge_key`. No-op if either is missing."""
    if not merge_key or not merge_cols:
        return
    headers = {c.value: c.column for c in ws[header_row]}
    if merge_key not in headers:
        return
    key_col = headers[merge_key]
    cols = [headers[h] for h in merge_cols if h in headers]
    if not cols:
        return
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r, last = header_row + 1, ws.max_row
    while r <= last:
        cur = ws.cell(row=r, column=key_col).value
        end = r
        while end + 1 <= last and ws.cell(row=end + 1, column=key_col).value == cur:
            end += 1
        if end > r:  # only multi-row groups
            for col in cols:
                ws.merge_cells(start_row=r, start_column=col, end_row=end, end_column=col)
                ws.cell(row=r, column=col).alignment = center
        r = end + 1


def generate_excel(
        rows: list[dict],
        sheet_name: str = "Report",
        merge_key: str | None = None,
        merge_cols: tuple[str, ...] | None = None,
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # Excel sheet name max 31 chars

    if not rows:
        ws.append(["No data available"])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    headers = list(rows[0].keys())

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Write headers
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Reusable styles — creating fresh style objects per cell is the main slowdown on
    # large sheets, so build them once and share. Column widths are tracked while writing
    # to avoid a second full-sheet pass.
    even_fill = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")
    odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    data_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    col_max = [len(str(h)) for h in headers]

    for row_idx, row in enumerate(rows, 2):
        fill = even_fill if row_idx % 2 == 0 else odd_fill
        for col_idx, header in enumerate(headers, 1):
            value = row.get(header, "")
            # Convert None to empty string
            if value is None:
                value = ""
            # Client requirement: all text values shown uppercase (numbers/dates untouched)
            if isinstance(value, str):
                value = value.upper()
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = data_align
            cell.border = thin_border
            cell.fill = fill

            n = len(value) if isinstance(value, str) else len(str(value))
            if n > col_max[col_idx - 1]:
                col_max[col_idx - 1] = n

    # Auto column width (computed above while writing)
    for i in range(len(headers)):
        ws.column_dimensions[get_column_letter(i + 1)].width = min(col_max[i] + 4, 45)

    # Row height for header
    ws.row_dimensions[1].height = 25

    # Freeze header row
    ws.freeze_panes = "A2"

    # Cell-merging is intentionally disabled: order-level values repeat on each line row
    # instead of being visually merged. This keeps output identical at any size, fast to
    # generate, and friendly to Excel sort/filter/pivot. (merge_key/merge_cols kept for
    # backward compatibility but no longer applied.)

    # Add generated timestamp in a metadata sheet
    meta_ws = wb.create_sheet("Info")
    meta_ws.append(["Generated At", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    meta_ws.append(["Total Rows", len(rows)])
    meta_ws.append(["Sheet", sheet_name])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── PDF ────────────────────────────────────────────────────────────────────────

def generate_pdf(
        rows: list[dict], title: str = "Report", col_weights: dict[str, float] | None = None,
) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        rightMargin=1 * cm,
        leftMargin=1 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    styles = getSampleStyleSheet()
    subtitle_style = ParagraphStyle(
        "Subtitle",
        parent=styles["Normal"],
        fontSize=8,
        textColor=colors.HexColor("#666666"),
        alignment=TA_CENTER,
    )
    cell_style = ParagraphStyle(
        "Cell",
        parent=styles["Normal"],
        fontSize=6.5,
        leading=8,
    )

    elements = []

    # Title + timestamp
    elements.append(Paragraph(title, styles["Title"]))
    elements.append(Paragraph(
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        subtitle_style,
    ))
    elements.append(Spacer(1, 0.5 * cm))

    if not rows:
        elements.append(Paragraph("No data available for the selected filters.", styles["Normal"]))
        doc.build(elements)
        return buf.getvalue()

    headers = list(rows[0].keys())
    page_width = landscape(A4)[0] - 2 * cm

    if col_weights:
        w = [col_weights.get(h, 1.0) for h in headers]
        total_w = sum(w)
        col_widths = [page_width * x / total_w for x in w]
    else:
        col_widths = [page_width / len(headers)] * len(headers)

    table_data = [[Paragraph(f"<b>{h}</b>", cell_style) for h in headers]]
    for row in rows:
        # Client requirement: all cell values shown uppercase (digits/dates unaffected)
        table_data.append([
            Paragraph(str(row.get(h, "") or "").upper().replace("\n", "<br/>"), cell_style)
            for h in headers
        ])

    table = Table(table_data, colWidths=col_widths, repeatRows=1)

    table.setStyle(TableStyle([
        # Header background
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        # Alternating row backgrounds
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
            colors.white,
            colors.HexColor("#EBF3FB"),
        ]),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("ALIGN", (0, 1), (-1, -1), "LEFT"),
        # Grid
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))

    elements.append(table)

    # Footer with page numbers
    def add_page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#999999"))
        page_text = f"Page {doc.page}"
        canvas.drawRightString(
            landscape(A4)[0] - 1 * cm,
            0.75 * cm,
            page_text,
        )
        canvas.drawString(
            1 * cm,
            0.75 * cm,
            f"{title} — {datetime.now().strftime('%Y-%m-%d')}",
        )
        canvas.restoreState()

    doc.build(elements, onFirstPage=add_page_number, onLaterPages=add_page_number)
    return buf.getvalue()


INK = colors.HexColor("#1f2933")
SUBTLE = colors.HexColor("#52606d")
MUTED = colors.HexColor("#9aa5b1")
LINE = colors.HexColor("#e4e7eb")
ZEBRA = colors.HexColor("#f7f9fb")


def _money(v):
    try:
        return f"{float(v or 0):,.2f}"
    except (TypeError, ValueError):
        return "0.00"


def _date(iso, fmt="%d %b %Y"):
    if not iso:
        return "-"
    try:
        return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(
            timezone.utc).strftime(fmt)
    except Exception:
        return str(iso)


def _styles():
    ss = getSampleStyleSheet()
    return {
        "brand": ParagraphStyle("brand", parent=ss["Normal"], fontName="Helvetica-Bold",
                                fontSize=18, textColor=INK, leading=22),
        "brand_sub": ParagraphStyle("brand_sub", parent=ss["Normal"], fontSize=8,
                                    textColor=MUTED, leading=11),
        "doc_title": ParagraphStyle("doc_title", parent=ss["Normal"], fontName="Helvetica-Bold",
                                    fontSize=20, textColor=INK, alignment=TA_RIGHT, leading=24),
        "meta": ParagraphStyle("meta", parent=ss["Normal"], fontSize=8.5, textColor=SUBTLE,
                                alignment=TA_RIGHT, leading=13),
        "lbl": ParagraphStyle("lbl", parent=ss["Normal"], fontSize=7.5, textColor=MUTED,
                              leading=12, spaceAfter=3),
        "name": ParagraphStyle("name", parent=ss["Normal"], fontName="Helvetica-Bold",
                               fontSize=10, textColor=INK, leading=14),
        "line": ParagraphStyle("line", parent=ss["Normal"], fontSize=9, textColor=SUBTLE,
                               leading=13),
        "th": ParagraphStyle("th", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=8,
                             textColor=colors.white, leading=10),
        "th_r": ParagraphStyle("th_r", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=8,
                               textColor=colors.white, leading=10, alignment=TA_RIGHT),
        "cell": ParagraphStyle("cell", parent=ss["Normal"], fontSize=9, textColor=INK, leading=12),
        "cell_meta": ParagraphStyle("cell_meta", parent=ss["Normal"], fontSize=7.5,
                                    textColor=MUTED, leading=10),
        "cell_r": ParagraphStyle("cell_r", parent=ss["Normal"], fontSize=9, textColor=INK,
                                 leading=12, alignment=TA_RIGHT),
        "foot": ParagraphStyle("foot", parent=ss["Normal"], fontSize=7.5, textColor=MUTED,
                               alignment=TA_LEFT, leading=11),
    }


def build_sales_order_pdf(order: dict) -> bytes:
    s = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=16 * mm, rightMargin=16 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
        title=f"Sales Order {order.get('orderNumber', '')}",
    )
    avail = doc.width
    story = []

    # ── Header: brand left, SALES ORDER + meta right ──
    ts = order.get("statusTimestamps") or {}
    meta = (
        f"<b>{order.get('orderNumber','')}</b><br/>"
        f"Date: {_date(ts.get('billedAt') or order.get('createdAt'))}<br/>"
        f"Order date: {_date(order.get('createdAt'))}<br/>"
        f"Status: {str(order.get('status','')).upper()}"
    )
    brand_cell = [
        Paragraph(order.get("tenantName") or "Sales Order", s["brand"]),
        Paragraph(f"Distributed via {order.get('distributorName') or '-'}", s["brand_sub"]),
    ]
    head = Table([[brand_cell,
                   [Paragraph("SALES ORDER", s["doc_title"]), Spacer(1, 4), Paragraph(meta, s["meta"])]]],
                 colWidths=[avail * 0.55, avail * 0.45])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                              ("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story += [head, Spacer(1, 8),
              HRFlowable(width="100%", thickness=1.4, color=INK, spaceAfter=14)]

    # ── Parties: billed-to / from ──
    addr = order.get("shopAddress") or {}
    addr_line = ", ".join(p for p in [addr.get("city"), addr.get("state_name"),
                                      addr.get("pincode")] if p)
    billed = [
        Paragraph("BILLED TO", s["lbl"]),
        Paragraph(order.get("shopName") or "-", s["name"]),
        Paragraph("<br/>".join(filter(None, [addr.get("line1"), addr.get("line2"), addr_line])),
                  s["line"]),
        Paragraph(f"Contact: {order.get('shopContactPerson') or '-'} &middot; "
                  f"{order.get('shopPhone') or '-'}", s["line"]),
    ]
    frm = [
        Paragraph("FROM / DISTRIBUTOR", s["lbl"]),
        Paragraph(order.get("distributorName") or "-", s["name"]),
        Paragraph(order.get("distributorPhone") or "-", s["line"]),
        Paragraph(f"Executive: {order.get('assignedExecutiveName') or '-'}", s["line"]),
    ]
    parties = Table([[billed, frm]], colWidths=[avail * 0.55, avail * 0.45])
    parties.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story += [parties, Spacer(1, 18)]

    # ── Line items ──
    header = [Paragraph("#", s["th"]), Paragraph("ITEM", s["th"]),
              Paragraph("QTY", s["th_r"]), Paragraph("UNIT PRICE", s["th_r"]),
              Paragraph("AMOUNT", s["th_r"])]
    rows = [header]
    for i, it in enumerate(order.get("items", []), start=1):
        name_cell = [Paragraph(str(it.get("productName") or "-"), s["cell"])]
        if it.get("setTypeName"):
            name_cell.append(Paragraph(str(it["setTypeName"]), s["cell_meta"]))
        rows.append([
            Paragraph(str(i), s["cell"]), name_cell,
            Paragraph(str(it.get("count", 0)), s["cell_r"]),
            Paragraph(_money(it.get("unitPrice")), s["cell_r"]),
            Paragraph(_money(it.get("totalPrice")), s["cell_r"]),
        ])
    cw = [avail * 0.06, avail * 0.50, avail * 0.10, avail * 0.16, avail * 0.18]
    items = Table(rows, colWidths=cw, repeatRows=1)
    st = [
        ("BACKGROUND", (0, 0), (-1, 0), INK),
        ("TOPPADDING", (0, 0), (-1, 0), 7), ("BOTTOMPADDING", (0, 0), (-1, 0), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 1), (-1, -1), 7), ("BOTTOMPADDING", (0, 1), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 1), (-1, -1), 0.5, LINE),
    ]
    for r in range(2, len(rows), 2):
        st.append(("BACKGROUND", (0, r), (-1, r), ZEBRA))
    items.setStyle(TableStyle(st))
    story += [items, Spacer(1, 14)]

    # ── Totals (right aligned block) ──
    disc_lbl = "Discount"
    if order.get("discountPercent"):
        disc_lbl = f"Discount ({order.get('discountPercent')}%)"
    tot_rows = [
        [Paragraph("Subtotal", s["line"]), Paragraph(_money(order.get("subtotal")), s["cell_r"])],
        [Paragraph(disc_lbl, s["line"]),
         Paragraph(f"- {_money(order.get('discountAmount'))}", s["cell_r"])],
        [Paragraph("<b>Total</b>", ParagraphStyle("g", parent=s["line"], fontSize=12,
                                                  textColor=INK)),
         Paragraph(f"<b>INR {_money(order.get('totalAmount'))}</b>",
                   ParagraphStyle("gr", parent=s["cell_r"], fontSize=12))],
    ]
    totals = Table(tot_rows, colWidths=[avail * 0.22, avail * 0.20])
    totals.setStyle(TableStyle([
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LINEABOVE", (0, 2), (-1, 2), 1.4, INK), ("TOPPADDING", (0, 2), (-1, 2), 9),
    ]))
    wrap = Table([[totals]], colWidths=[avail])
    wrap.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "RIGHT"),
                             ("LEFTPADDING", (0, 0), (-1, -1), 0),
                             ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story.append(wrap)

    # ── Notes ──
    if (order.get("notes") or "").strip():
        story += [Spacer(1, 18),
                  Paragraph("NOTES", s["lbl"]),
                  Paragraph(order["notes"].strip(), s["line"])]

    # ── Footer ──
    story += [Spacer(1, 28),
              HRFlowable(width="100%", thickness=0.5, color=LINE, spaceAfter=8),
              Paragraph(f"System-generated sales order for {order.get('orderNumber','')}. "
                        f"No signature required.", s["foot"])]

    doc.build(story)
    return buf.getvalue()